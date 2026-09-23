"""Render FinalReport.md to a formatted .docx.

Handles the Markdown subset the report actually uses: ATX headings, pipe tables,
fenced code blocks (including the ASCII figures), image embeds, ordered lists,
figure/table legends, and inline bold / italic / code spans.

    python backend/scripts/build_docx.py [output.docx]
"""

from __future__ import annotations

import pathlib
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parents[2]
SOURCE = ROOT / "FinalReport.md"

# Palette drawn from the product's own identity (ink / amber / moss).
INK = RGBColor(0x1A, 0x16, 0x14)
AMBER = RGBColor(0xA8, 0x6A, 0x12)
MUTED = RGBColor(0x5A, 0x52, 0x4C)
RULE = "C9BFB4"
HEAD_BG = "F3EEE8"

BODY_FONT = "Calibri"
HEAD_FONT = "Georgia"
MONO_FONT = "Consolas"

MAX_IMG_W = Cm(15.0)
MAX_IMG_H = Cm(13.0)

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`)", re.S)


# ---------------------------------------------------------------- low level


def _shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def _border(paragraph, edge: str, size: int, color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    borders = pPr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        pPr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), color)
    borders.append(el)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, sep, end):
        run._r.append(node)


def add_runs(paragraph, text: str, *, size: Pt | None = None,
             color: RGBColor | None = None, italic_all: bool = False,
             bold_all: bool = False):
    """Add `text` to `paragraph`, honouring **bold**, *italic* and `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        bold = italic = code = False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        if (bold or italic) and INLINE.search(part):
            add_runs(paragraph, part, size=size, color=color,
                     italic_all=italic or italic_all, bold_all=bold or bold_all)
            continue
        run = paragraph.add_run(part)
        run.bold = bold or bold_all
        run.italic = italic or italic_all
        if code:
            run.font.name = MONO_FONT
            run.font.size = Pt((size or Pt(10.5)).pt - 0.5)
        elif size:
            run.font.size = size
        if color:
            run.font.color.rgb = color
    return paragraph


# ---------------------------------------------------------------- styles


def build_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.15

    for name, size, space_before, color in (
        ("Heading 1", 19, 22, INK),
        ("Heading 2", 13, 15, INK),
        ("Heading 3", 11.5, 12, MUTED),
    ):
        st = doc.styles[name]
        st.font.name = HEAD_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)


def running_head(section, title: str) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(p, title, size=Pt(8), color=MUTED)
    _border(p, "bottom", 4, RULE)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(fp, " PAGE ")
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


# ---------------------------------------------------------------- blocks


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    add_runs(p, text, size=Pt(9), color=MUTED)


def add_code(doc: Document, lines: list[str]) -> None:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(10)
    pf.left_indent = Cm(0.3)
    pf.line_spacing = 1.0
    pf.keep_together = True
    _shade(p._p.get_or_add_pPr(), "F7F4F0")
    _border(p, "left", 12, "D8CFC4")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = MONO_FONT
        run.font.size = Pt(7.6)
        run.font.color.rgb = INK
        if i < len(lines) - 1:
            run.add_break()


def add_image(doc: Document, path: pathlib.Path) -> None:
    with Image.open(path) as im:
        w, h = im.size
    width = MAX_IMG_W
    if Cm(width.cm * h / w) > MAX_IMG_H:
        width = Cm(MAX_IMG_H.cm * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=width)


TABLE_W = Cm(16.0)


def _column_widths(rows: list[list[str]]) -> list[Cm]:
    """Size columns by their longest cell, so prose columns get the room."""
    cols = len(rows[0])
    demand = []
    for c in range(cols):
        longest = max(len(re.sub(r"[*`]", "", r[c])) for r in rows)
        header = len(re.sub(r"[*`]", "", rows[0][c]))
        demand.append(max(header * 1.35 + 2.0, min(float(longest), 46.0)))
    total = sum(demand)
    widths = [TABLE_W.cm * d / total for d in demand]
    # keep every column usable, then re-normalise
    widths = [max(1.6, w) for w in widths]
    scale = TABLE_W.cm / sum(widths)
    return [Cm(w * scale) for w in widths]


def add_table(doc: Document, rows: list[list[str]], aligns: list[str]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = _column_widths(rows)
    table._tbl.tblPr.append(
        OxmlElement("w:tblLayout")
    )
    table._tbl.tblPr[-1].set(qn("w:type"), "fixed")

    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.0
            if c < len(aligns) and aligns[c] == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cell.width = widths[c]
            if r == 0:
                _shade(cell._tc.get_or_add_tcPr(), HEAD_BG)
                run = p.add_run(re.sub(r"\*\*", "", cell_text))
                run.bold = True
                run.font.size = Pt(9)
            else:
                add_runs(p, cell_text, size=Pt(9))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def title_page(doc: Document, meta: dict[str, str]) -> None:
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("NextTrack")
    run.font.name = HEAD_FONT
    run.font.size = Pt(40)
    run.bold = True
    run.font.color.rgb = INK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("A Stateless Music Recommendation API")
    run.font.name = HEAD_FONT
    run.font.size = Pt(16)
    run.font.color.rgb = AMBER

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(20)
    _border(rule, "bottom", 6, RULE)

    for line, bold in (
        ("CM3070 Final Year Project — Final Project Report", True),
        ("University of London International Programmes", False),
        ("BSc Computer Science and Related Subjects", False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(11)
        if not bold:
            run.font.color.rgb = MUTED

    doc.add_paragraph()
    for label, value in meta.items():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(f"{label}  ")
        run.font.size = Pt(10)
        run.font.color.rgb = MUTED
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.bold = True

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def toc_page(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run("Contents")
    run.font.name = HEAD_FONT
    run.font.size = Pt(19)
    run.bold = True
    run.font.color.rgb = INK
    p.paragraph_format.space_after = Pt(4)
    _border(p, "bottom", 6, RULE)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    _field(doc.add_paragraph(), r' TOC \o "1-2" \h \z \u ')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- parser


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md: str, doc: Document) -> None:
    lines = md.split("\n")
    i = 0
    total = len(lines)
    in_references = False

    while i < total:
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # fenced code / ASCII figure
        if stripped.startswith("```"):
            i += 1
            block: list[str] = []
            while i < total and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip())
                i += 1
            i += 1
            while block and not block[0].strip():
                block.pop(0)
            while block and not block[-1].strip():
                block.pop()
            if block:
                add_code(doc, block)
            continue

        # image
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", stripped)
        if m:
            path = ROOT / m.group(1)
            if path.exists():
                add_image(doc, path)
            i += 1
            continue

        # headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:].strip()
            in_references = heading.startswith("References")
            if heading.startswith(("Chapter", "References", "Appendix")):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(heading, level=1)
            _border(h, "bottom", 6, RULE)
            i += 1
            continue
        if stripped.startswith("# "):
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < total and re.match(
            r"^\|[\s:|-]+\|$", lines[i + 1].strip()
        ):
            header = split_row(stripped)
            aligns = [
                "right" if c.strip().endswith(":") else "left"
                for c in split_row(lines[i + 1])
            ]
            rows = [header]
            i += 2
            while i < total and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            width = len(header)
            rows = [(r + [""] * width)[:width] for r in rows]
            add_table(doc, rows, aligns)
            continue

        # ordered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            body = [m.group(2)]
            i += 1
            while i < total and lines[i].startswith("   ") and lines[i].strip():
                body.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, " ".join(body))
            continue

        # paragraph (gather the wrapped block)
        block = [stripped]
        i += 1
        while i < total:
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", "```", "!["))
                    or re.match(r"^\d+\.\s", nxt) or nxt == "---"):
                break
            block.append(nxt)
            i += 1
        text = " ".join(block)

        if re.match(r"^\*{1,2}(Figure|Table)\s+\d+", text):
            add_caption(doc, text)
        else:
            para = doc.add_paragraph()
            if in_references:
                pf = para.paragraph_format
                pf.left_indent = Cm(0.8)
                pf.first_line_indent = Cm(-0.8)
                pf.space_after = Pt(7)
            add_runs(para, text)


def main() -> int:
    out = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "NextTrack_Final_Report.docx"
    md = SOURCE.read_text(encoding="utf-8")

    doc = Document()
    build_styles(doc)
    page_setup(doc)

    title_page(doc, {
        "Author": "Soe Ming Wei, Glenn",
        "Student number": "230657168",
        "Supervisor": "Yeo Sze Wee",
        "Project template": "7.2 — Project Idea 2: NextTrack",
        "Code repository": "github.com/glennysmw/NextTrack",
    })
    toc_page(doc)

    body = md[md.index("## Chapter 1"):]
    convert(body, doc)

    running_head(doc.sections[0], "NextTrack · CM3070 Final Project Report")
    doc.save(out)
    print(f"wrote {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
